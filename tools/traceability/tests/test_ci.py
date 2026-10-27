"""Security-sensitive, fail-closed CI evidence tests (no live API required)."""
import base64
import io
import zipfile

import httpx
import pytest

from openpyxl import load_workbook

from samd_traceability_mcp.ci import (
    extract_junit_zip, enrich_matrix_with_ci, parse_junit_xml,
    validated_artifact_redirect,
)
from samd_traceability_mcp.excel import make_xlsx_response
from samd_traceability_mcp.matrix import build_traceability_matrix


def junit(*entries):
    items = []
    for name, state in entries:
        outcome = {"Passed": "", "Failed": "<failure/>", "Skipped": "<skipped/>"}[state]
        items.append(
            '<testcase classname="tests.test_example" name="' + name + '">' + outcome + '</testcase>'
        )
    return ('<testsuites><testsuite>' + "".join(items) + '</testsuite></testsuites>').encode()


def report():
    return {
        "repository": "example-org/example-repo",
        "ref": "test/notion-e2e-demo-fixtures",
        "notion_data_source_id": "00000000-0000-4000-8000-000000000001",
        "requirements_count": 2, "linked_count": 1, "unlinked_count": 1,
        "requirements": [
            {"id": "REQ-DEM-001", "title": "Readings", "url": "https://www.notion.so/example",
             "test_files": ["tests/test_example.py"],
             "test_cases": [{"test_file": "tests/test_example.py",
                             "test_name": "test_readings", "line": 2}]},
            {"id": "REQ-DEM-002", "title": "Trends", "url": "https://www.notion.so/other",
             "test_files": [], "test_cases": []},
        ],
        "references_to_unknown_requirement_ids": [],
    }


def test_parse_test_function_identity_and_outcome():
    cases = parse_junit_xml(junit(
        ("test_readings", "Passed"), ("test_other", "Failed"),
        ("test_skip", "Skipped")))
    assert cases == {
        ("tests/test_example.py", "test_readings"): "Passed",
        ("tests/test_example.py", "test_other"): "Failed",
        ("tests/test_example.py", "test_skip"): "Skipped",
    }


def test_multiple_parameters_aggregate_failures():
    cases = parse_junit_xml(junit(
        ("test_readings[a]", "Passed"), ("test_readings[b]", "Failed")))
    assert cases[("tests/test_example.py", "test_readings")] == "Failed"


def test_wrong_module_and_non_test_names_do_not_link():
    xml = b'<testsuite><testcase classname="elsewhere.other" name="test_readings"/><testcase classname="tests.test_example" name="helper"/></testsuite>'
    cases = parse_junit_xml(xml)
    matrix = enrich_matrix_with_ci(build_traceability_matrix(report()), {
        "status": "Available", "run_url": "https://github.com/example/actions/runs/1",
        "sha": "a" * 40, "cases": cases})
    assert matrix["rows"][0]["Execution Status"] == "Unknown"
    assert matrix["rows"][1]["Execution Status"] == "Unknown"


def test_unavailable_ci_never_infers_pass():
    matrix = enrich_matrix_with_ci(build_traceability_matrix(report()), {
        "status": "Unavailable", "reason": "No matching commit", "cases": {}})
    assert all(r["Execution Status"] == "Unknown" for r in matrix["rows"])
    assert matrix["ci_evidence"]["matched_test_cases"] == 0


def test_exact_test_case_evidence_and_xlsx_link():
    matrix = enrich_matrix_with_ci(build_traceability_matrix(report()), {
        "status": "Available",
        "run_url": "https://github.com/example-org/example-repo/actions/runs/17",
        "sha": "a" * 40,
        "cases": {("tests/test_example.py", "test_readings"): "Passed"},
    })
    assert matrix["rows"][0]["Execution Status"] == "Passed"
    assert matrix["rows"][0]["CI Commit SHA"] == "a" * 40
    assert matrix["rows"][1]["Execution Status"] == "Unknown"
    raw = base64.b64decode(make_xlsx_response(matrix)["content_base64"])
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Traceability Matrix"]
    headers = {c.value: c.column for c in ws[1]}
    assert ws.cell(2, headers["Execution Status"]).value == "Passed"
    assert ws.cell(2, headers["CI Run URL"]).hyperlink.target.endswith("/17")


def test_invalid_or_oversized_junit_rejected():
    import pytest
    with pytest.raises(ValueError, match="root"):
        parse_junit_xml(b"<other/>")
    with pytest.raises(ValueError, match="size limit"):
        parse_junit_xml(b"x" * 3_000_001)


def test_named_zip_artifact():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("pytest-results.xml", junit(("test_readings", "Passed")))
    assert extract_junit_zip(buf.getvalue())[("tests/test_example.py", "test_readings")] == "Passed"


def test_artifact_download_accepts_github_redirect_without_forwarding_token():
    response = httpx.Response(
        302,
        request=httpx.Request("GET", "https://api.github.com/repos/example/repo/actions/artifacts/5/zip"),
        headers={"Location": "https://objects.githubusercontent.com/signed-artifact"},
    )
    assert validated_artifact_redirect(response) == (
        "https://objects.githubusercontent.com/signed-artifact"
    )


def test_artifact_download_rejects_non_storage_redirect_and_real_http_errors():
    request = httpx.Request("GET", "https://api.github.com/repos/example/repo/actions/artifacts/5/zip")
    with pytest.raises(ValueError, match="Invalid GitHub artifact redirect"):
        validated_artifact_redirect(httpx.Response(
            302, request=request, headers={"Location": "https://malicious.example/collect"},
        ))
    with pytest.raises(httpx.HTTPStatusError):
        validated_artifact_redirect(httpx.Response(403, request=request))
