import base64
import io

from openpyxl import load_workbook

from samd_traceability_mcp.excel import make_xlsx_response
from samd_traceability_mcp.matrix import build_traceability_matrix
from samd_traceability_mcp.notion import analyze_notion_links, parse_notion_requirements


def notion_page(identifier, title):
    def text(value):
        return [{"plain_text": value}]
    return {
        "url": f"https://www.notion.so/{identifier}",
        "properties": {
            "Name": {"type": "title", "title": text(title)},
            "Requirement ID": {"type": "rich_text", "rich_text": text(identifier)},
            "Description": {"type": "rich_text", "rich_text": text("Synthetic sample display requirement")},
            "Priority": {"type": "select", "select": {"name": "High"}},
            "Status": {"type": "select", "select": {"name": "Approved"}},
            "Release": {"type": "select", "select": {"name": "R1"}},
        },
    }


def synthetic_matrix():
    reqs = parse_notion_requirements([
        notion_page("REQ-DEM-001", "Latest sample measurement"),
        notion_page("REQ-DEM-002", "Sample history"),
    ])
    source = '''
def test_sample_display_demo():
    """Requirement: REQ-DEM-001"""
    assert True

def test_sample_display_secondary_demo():
    """Requirement: REQ-DEM-001"""
    assert True

def test_dangling_demo():
    """Requirement: REQ-OLD-999"""
    assert True
'''
    report = analyze_notion_links(
        reqs,
        {"tests/test_sample.py": source},
        "example-org/example-repo",
    )
    report["ref"] = "test/notion-e2e-demo-fixtures"
    report["notion_data_source_id"] = "00000000-0000-4000-8000-000000000001"
    return build_traceability_matrix(report)


def test_case_level_link_and_metadata():
    matrix = synthetic_matrix()
    assert matrix["linked_count"] == 1
    assert matrix["unlinked_count"] == 1
    assert matrix["unknown_reference_count"] == 1
    assert matrix["row_count"] == 4
    assert {r["Test Case"] for r in matrix["rows"][:2]} == {
        "test_sample_display_demo", "test_sample_display_secondary_demo",
    }
    assert matrix["rows"][0]["GitHub Test URL"].endswith("#L2")
    assert matrix["rows"][0]["Release"] == "R1"
    assert matrix["rows"][0]["Requirement Status"] == "Approved"
    assert matrix["rows"][0]["Priority"] == "High"
    assert matrix["rows"][2]["Link Status"] == "Unlinked"
    assert matrix["rows"][3]["Link Status"] == "Unknown requirement ID"


def test_excel_export_openable_with_filters_and_clickable_source_links():
    matrix = synthetic_matrix()
    result = make_xlsx_response(matrix)
    raw = base64.b64decode(result["content_base64"])
    assert raw[:2] == b"PK"
    book = load_workbook(io.BytesIO(raw))
    assert book.sheetnames == ["Overview", "Traceability Matrix"]
    sheet = book["Traceability Matrix"]
    assert len(sheet.tables) == 1
    assert sheet.freeze_panes == "F2"
    headers = [cell.value for cell in sheet[1]]
    assert "Test Case" in headers
    row = list(sheet.iter_rows(min_row=2, max_row=2))[0]
    cells = dict(zip(headers, row))
    assert cells["Test Case"].value == "test_sample_display_demo"
    assert cells["Notion URL"].hyperlink.target.startswith("https://www.notion.so/")
    assert "#L2" in cells["GitHub Test URL"].hyperlink.target
    assert cells["Execution Status"].value == "Unknown"
    assert book["Overview"]["B11"].value == matrix["execution_status_policy"]


def test_excel_strings_cannot_be_interpreted_as_formulas():
    matrix = synthetic_matrix()
    matrix["rows"][0]["Requirement Title"] = "=1+1"
    raw = base64.b64decode(make_xlsx_response(matrix)["content_base64"])
    sheet = load_workbook(io.BytesIO(raw))["Traceability Matrix"]
    idx = [c.value for c in sheet[1]].index("Requirement Title") + 1
    assert sheet.cell(2, idx).data_type != "f"
    assert sheet.cell(2, idx).value == "'=1+1"
