"""Professional, in-memory Excel export for Notion/GitHub traceability evidence."""
from __future__ import annotations

import base64
import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

GRAPHITE = "202C37"
TEAL = "147D84"
PALE_TEAL = "EAF6F5"
PALE_ORANGE = "FFF2DF"
PALE_RED = "FCEBEC"
WHITE = "FFFFFF"
MUTED = "5B6877"

VISIBLE_COLUMNS = [
    ("Requirement ID", 20),
    ("Requirement Title", 43),
    ("Release", 12),
    ("Requirement Status", 19),
    ("Priority", 13),
    ("Test Case", 49),
    ("GitHub Test File", 45),
    ("Link Status", 23),
    ("Execution Status", 22),
    ("CI Run URL", 20),
    ("CI Commit SHA", 46),
    ("Evidence Level", 27),
    ("Notion URL", 18),
    ("GitHub Test URL", 18),
    ("Description", 68),
    ("Record Type", 26),
]


def _safe_excel_text(value: object) -> str:
    text = str(value if value is not None else "")
    # Stop external spreadsheet data being evaluated as an Excel formula.
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def render_xlsx_bytes(matrix: dict) -> bytes:
    """Produce a styled Excel workbook without writing to the filesystem."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Overview"
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:H2")
    title = summary["A1"]
    title.value = "Requirements Traceability Matrix"
    title.fill = PatternFill("solid", fgColor=GRAPHITE)
    title.font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    title.alignment = Alignment(vertical="center", indent=1)
    for row in summary["A1:H2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=GRAPHITE)
    summary.row_dimensions[1].height = 24
    summary.row_dimensions[2].height = 24
    summary["A4"] = "REQUIREMENTS"
    summary["C4"] = "LINKED"
    summary["E4"] = "UNLINKED"
    summary["G4"] = "UNKNOWN REFERENCES"
    values = [
        (1, matrix["requirements_count"]),
        (3, matrix["linked_count"]),
        (5, matrix["unlinked_count"]),
        (7, matrix["unknown_reference_count"]),
    ]
    for col, val in values:
        lab = summary.cell(4, col)
        lab.font = Font(name="Aptos", size=10, bold=True, color=MUTED)
        cell = summary.cell(5, col)
        cell.value = val
        cell.font = Font(name="Aptos Display", size=23, color=TEAL, bold=True)
    summary["A8"] = "Scope"
    summary["B8"] = "Traceability reflects the evidence available at export time"
    summary["A9"] = "Repository"
    summary["B9"] = matrix["repository"]
    summary["A10"] = "Git ref"
    summary["B10"] = matrix["ref"]
    summary["A11"] = "Execution evidence"
    summary["B11"] = matrix.get("execution_status_policy", "Not collected; every test execution status is Unknown")
    ci = matrix.get("ci_evidence", {})
    if ci:
        summary["A16"] = "CI evidence"
        summary["B16"] = ci.get("status", "Unavailable")
        summary["A17"] = "Matched test cases"
        summary["B17"] = ci.get("matched_test_cases", 0)
        summary["A18"] = "CI run URL"
        summary["B18"] = ci.get("run_url", "")
        if ci.get("run_url", "").startswith("https://github.com/"):
            summary["B18"].hyperlink = ci["run_url"]
        summary["A19"] = "CI evidence note"
        summary["B19"] = ci.get("reason", "")
    summary["A13"] = "IMPORTANT"
    summary["B13"] = "A traceability marker is not proof of a passing or adequate test."
    summary["A13"].font = Font(bold=True, color=GRAPHITE)
    summary["B13"].font = Font(color="9A4D00")
    summary["A15"] = "Data source"
    summary["B15"] = "Notion requirements and GitHub test docstrings"
    for col in "ABCDEFGH":
        summary.column_dimensions[col].width = 17 if col != "B" else 56
    for row in (8, 9, 10, 11, 13, 15):
        summary.cell(row, 1).font = Font(name="Aptos", bold=True, color=GRAPHITE)
        summary.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="center")
    summary.column_dimensions["B"].width = 70
    summary.row_dimensions[11].height = 30
    summary.row_dimensions[13].height = 30

    ws = wb.create_sheet("Traceability Matrix")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "F2"
    for index, (header, width) in enumerate(VISIBLE_COLUMNS, 1):
        cell = ws.cell(1, index, header)
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=GRAPHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 34

    for idx, row in enumerate(matrix["rows"], 2):
        for col, (key, _width) in enumerate(VISIBLE_COLUMNS, 1):
            value = _safe_excel_text(row.get(key, ""))
            cell = ws.cell(idx, col, value)
            cell.font = Font(name="Aptos", size=10, color=GRAPHITE)
            cell.alignment = Alignment(vertical="center", wrap_text=key in ("Requirement Title", "Test Case"))
            if key in ("Notion URL", "GitHub Test URL", "CI Run URL") and value.startswith("https://"):
                cell.hyperlink = value
                cell.value = "Open Notion" if key == "Notion URL" else "Open CI Run" if key == "CI Run URL" else "Open GitHub"
                cell.font = Font(name="Aptos", size=10, color=TEAL, underline="single")
            if key == "Link Status":
                cell.fill = PatternFill("solid", fgColor=PALE_TEAL if value == "Linked"
                                        else PALE_ORANGE if value == "Unlinked" else PALE_RED)
                cell.font = Font(name="Aptos", bold=True, color=GRAPHITE)
            if key == "Execution Status":
                cell.fill = PatternFill("solid", fgColor=PALE_TEAL if value == "Passed"
                                        else PALE_RED if value == "Failed" else PALE_ORANGE)
        ws.row_dimensions[idx].height = 32

    if matrix["rows"]:
        last_col = get_column_letter(len(VISIBLE_COLUMNS))
        tab = Table(displayName="RequirementsTraceability", ref=f"A1:{last_col}{len(matrix['rows']) + 1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                            showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:1"
    wb.active = 0
    buff = io.BytesIO()
    wb.save(buff)
    return buff.getvalue()


def make_xlsx_response(matrix: dict) -> dict:
    payload = render_xlsx_bytes(matrix)
    return {
        "file_name": "requirements_traceability_matrix.xlsx",
        "format": "xlsx",
        "encoding": "base64",
        "content_base64": base64.b64encode(payload).decode("ascii"),
        "bytes_count": len(payload),
        "requirements_count": matrix["requirements_count"],
        "linked_count": matrix["linked_count"],
        "unlinked_count": matrix["unlinked_count"],
        "unknown_reference_count": matrix["unknown_reference_count"],
        "execution_status_policy": matrix.get("execution_status_policy", "Unknown, no execution results were collected"),
        "ci_evidence": matrix.get("ci_evidence", {"status": "Not requested"}),
        "note": "Save decoded base64 locally. The MCP does not modify source systems.",
    }
